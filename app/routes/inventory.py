from flask import Blueprint, render_template, request, session, redirect, url_for, flash, send_file, jsonify, current_app
from ..database import get_db_connection
from datetime import datetime
import io
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..utils.csrf import csrf_required
import pyodbc
from ..utils.security import validate_file_upload, MAX_UPLOAD_SIZE, safe_error, limiter

inventory_bp = Blueprint('inventory', __name__)

@inventory_bp.route('/inventory', methods=['GET'])
def list_inventory():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    conn = get_db_connection()
    cursor = conn.cursor()

    # Get all inventory data
    cursor.execute("""
        SELECT 
            A.AssignmentID,
            DT.TypeName, DT.Brand, DT.Model,
            A.Serial,
            A.AssignmentDate,
            COALESCE(S.FirstName + ' ' + S.LastName, '') AS StaffName,
            W.Name AS WarehouseName,
            COALESCE(S.RUT, 'BODEGA') AS RUT,
            T.Name AS TerminalName,
            A.TerminalID,
            C.Name AS CompanyName,
            A.DeviceTypeID, A.StaffID, A.WarehouseID, A.AssetType, A.CompanyID,
            CASE WHEN A.WarehouseID IS NOT NULL THEN 'warehouse' ELSE 'staff' END AS HolderType,
            U.FullName AS AssignedBy,
            A.IsMobile, A.MobileCompany, A.MobileNumber, A.MobileSIM,
            A.IsTemporary, A.ExpectedReturnDate
        FROM Assignments A
        LEFT JOIN Staff S ON A.StaffID = S.StaffID
        LEFT JOIN Warehouses W ON A.WarehouseID = W.WarehouseID
        LEFT JOIN Terminals T ON A.TerminalID = T.TerminalID
        LEFT JOIN Companies C ON A.CompanyID = C.CompanyID
        LEFT JOIN DeviceTypes DT ON A.DeviceTypeID = DT.DeviceTypeID
        LEFT JOIN Users U ON A.UserID = U.UserID
        ORDER BY COALESCE(T.Name, 'ZZZ'), COALESCE(S.FirstName, 'ZZZ'), COALESCE(W.Name, 'ZZZ')
    """)
    
    inventory = cursor.fetchall()

    # Get lookups
    cursor.execute("SELECT TerminalID, Name FROM Terminals ORDER BY Name")
    terminals = cursor.fetchall()
    
    cursor.execute("SELECT WarehouseID, Name FROM Warehouses ORDER BY Name")
    warehouses = cursor.fetchall()
    
    cursor.execute("SELECT DeviceTypeID, TypeName, Brand, Model FROM DeviceTypes ORDER BY TypeName")
    device_types = cursor.fetchall()
    
    cursor.execute("SELECT StaffID, FirstName, LastName, RUT FROM Staff ORDER BY FirstName")
    staff = cursor.fetchall()
    
    cursor.execute("SELECT CompanyID, Name FROM Companies ORDER BY Name")
    companies = cursor.fetchall()

    conn.close()

    return render_template('inventory.html',
                          inventory=inventory,
                          terminals=terminals if terminals else [],
                          device_types=device_types if device_types else [],
                          staff=staff if staff else [],
                          warehouses=warehouses if warehouses else [],
                          companies=companies if companies else [],
                          now=datetime.now())

@inventory_bp.route('/inventory/edit', methods=['POST'])
@csrf_required
@limiter.limit("20 per minute")
def edit_inventory():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    if not session.get('permissions', {}).get('edit_inventory'):
        flash("No tiene permisos para editar inventario.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    assignment_id = request.form.get('assignment_id')
    device_type_id = request.form.get('device_type_id')
    serial = request.form.get('serial')
    staff_id = request.form.get('staff_id')
    warehouse_id = request.form.get('warehouse_id')
    asset_type = request.form.get('asset_type')
    company_id = request.form.get('company_id')
    date_val = datetime.now()

    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        terminal_id = request.form.get('terminal_id')
        is_mobile = 1 if request.form.get('is_mobile') == 'on' else 0
        mobile_company = request.form.get('mobile_company')
        mobile_number = request.form.get('mobile_number')
        mobile_sim = request.form.get('mobile_sim')

        cursor.execute("""
            EXEC sp_UpdateAssignment 
            @AssignmentID=?, @DeviceTypeID=?, @Serial=?, @Date=?, @StaffID=?, @WarehouseID=?,
            @AssetType=?, @CompanyID=?, @IsTemporary=0, @ExpectedReturnDate=NULL, @UserID=?,
            @IsMobile=?, @MobileCompany=?, @MobileNumber=?, @MobileSIM=?, @TerminalID=?
        """, (assignment_id, device_type_id, serial, date_val, staff_id or None, warehouse_id or None,
             asset_type, company_id, session.get('user_id'), is_mobile, mobile_company, mobile_number, mobile_sim, terminal_id or None))
        
        conn.commit()
        flash("Inventario actualizado correctamente.", "success")
    except Exception as e:
        flash("Error al actualizar el inventario. Contacte a soporte.", "danger")
    finally:
        conn.close()

    return redirect(url_for('inventory.list_inventory'))

@inventory_bp.route('/inventory/history/<serial>')
def view_history(serial):
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    if not session.get('permissions', {}).get('view_history'):
        return jsonify({"history": [], "error": "No tiene permisos para ver el historial."})

    if not serial or len(serial.strip()) == 0:
        return jsonify({"history": [], "error": "Serial requerido"})

    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        cursor.execute("""
            SELECT TOP 20
                HistoryID, Username, StaffRUT, TerminalName,
                AssignmentType, AssignmentModel, AssignmentSerial,
                AssignmentDate, Action, PreviousUser
            FROM AssignmentHistory WITH (NOLOCK)
            WHERE AssignmentSerial = ?
            ORDER BY AssignmentDate DESC
        """, (serial,))
        history = cursor.fetchall()
    except Exception as e:
        return jsonify({"history": [], "error": "Error al consultar el historial."})
    finally:
        conn.close()
    
    # Convert dates to string
    result = []
    for row in history:
        row_data = list(row)
        if row_data[7] and isinstance(row_data[7], datetime):
            row_data[7] = row_data[7].isoformat()
        result.append(row_data)
    
    return jsonify({"history": result})

@inventory_bp.route('/inventory/export')
def export_inventory():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))
    
    if not session.get('permissions', {}).get('export_inventory'):
        flash("No tiene permisos para exportar el inventario.", "danger")
        return redirect(url_for('inventory.list_inventory'))
    
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Border, Side

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT 
            COALESCE(T.Name, TW.Name) AS Terminal,
            C.Name AS Empresa,
            COALESCE(S.RUT, 'BODEGA') AS RUT,
            COALESCE(S.FirstName + ' ' + S.LastName, 'BODEGA: ' + W.Name) AS Funcionario,
            DT.TypeName AS Tipo, DT.Brand AS Marca, DT.Model AS Modelo,
            A.Serial AS Serie, A.AssetType AS Activo,
            CASE WHEN A.IsTemporary = 1 THEN 'Temporal' ELSE 'Indefinida' END AS TipoAsignacion,
            CONVERT(VARCHAR, A.ExpectedReturnDate, 103) AS FechaDevolucion,
            CASE WHEN A.IsMobile = 1 THEN 'Sí' ELSE 'No' END AS Movil,
            A.MobileCompany AS Compania, A.MobileNumber AS Linea, A.MobileSIM AS [NUMERO SERIE SIM],
            CONVERT(VARCHAR, A.AssignmentDate, 103) AS Fecha,
            U.FullName AS Usuario
        FROM Assignments A
        LEFT JOIN Staff S ON A.StaffID = S.StaffID
        LEFT JOIN Terminals T ON S.TerminalID = T.TerminalID
        LEFT JOIN Warehouses W ON A.WarehouseID = W.WarehouseID
        LEFT JOIN Terminals TW ON W.TerminalID = TW.TerminalID
        LEFT JOIN Companies C ON A.CompanyID = C.CompanyID
        LEFT JOIN DeviceTypes DT ON A.DeviceTypeID = DT.DeviceTypeID
        LEFT JOIN Users U ON A.UserID = U.UserID
        ORDER BY Terminal, Funcionario
    """)
    rows = cursor.fetchall()
    columns = [col[0] for col in cursor.description]
    conn.close()

    df = pd.DataFrame.from_records([list(r) for r in rows], columns=columns)

    # Create Excel with HGT styles
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Inventario HGT')
        ws = writer.sheets['Inventario HGT']

        # Styles
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='212A37', end_color='212A37', fill_type='solid')
        data_font = Font(name='Arial', size=10)
        border = Border(
            left=Side(style='thin', color='212A37'),
            right=Side(style='thin', color='212A37'),
            top=Side(style='thin', color='212A37'),
            bottom=Side(style='thin', color='212A37')
        )
        
        # Apply styles
        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=len(rows)+1), 1):
            for cell in row:
                cell.border = border
                if row_num == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                else:
                    cell.font = data_font
                    cell.fill = PatternFill(start_color='E8E8E8' if row_num % 2 == 0 else 'FFFFFF')
        
        # Auto-fit columns
        for col in ws.columns:
            max_length = 10
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 35)
    
    output.seek(0)

    # Find column indexes for conditional formatting
    tipo_col_idx = None
    fecha_dev_col_idx = None
    for idx, col in enumerate(columns, 1):
        if col == 'TipoAsignacion':
            tipo_col_idx = idx
        if col == 'FechaDevolucion':
            fecha_dev_col_idx = idx

    # Re-open to apply conditional coloring on Tipo column
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill as PF
    output.seek(0)
    wb = load_workbook(output)
    ws2 = wb.active
    temporal_fill = PF(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
    overdue_fill  = PF(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
    if tipo_col_idx:
        for row in ws2.iter_rows(min_row=2):
            tipo_cell = row[tipo_col_idx - 1]
            if tipo_cell.value == 'Temporal':
                tipo_cell.fill = temporal_fill
                if fecha_dev_col_idx:
                    row[fecha_dev_col_idx - 1].fill = temporal_fill
            # Check if date column has value and if the date is past
            if fecha_dev_col_idx and row[fecha_dev_col_idx - 1].value and tipo_cell.value == 'Temporal':
                from datetime import datetime as _dt
                try:
                    ret_date = _dt.strptime(str(row[fecha_dev_col_idx - 1].value), '%d/%m/%Y')
                    if ret_date < _dt.now():
                        tipo_cell.fill = overdue_fill
                        row[fecha_dev_col_idx - 1].fill = overdue_fill
                except Exception:
                    pass

    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)

    return send_file(
        final_output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'Inventario_HGT_{datetime.now().strftime("%d%m%Y")}.xlsx'
    )

@inventory_bp.route('/inventory/template')
def download_template():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT TerminalID, Name FROM Terminals ORDER BY Name")
    terminals = cursor.fetchall()

    cursor.execute("SELECT CompanyID, Name FROM Companies ORDER BY Name")
    companies = cursor.fetchall()

    cursor.execute("SELECT DeviceTypeID, TypeName, Brand, Model FROM DeviceTypes ORDER BY TypeName")
    device_types = cursor.fetchall()

    conn.close()

    columns = ['Terminal', 'RUT', 'Nombre', 'Apellido', 'AreaTrabajo', 'Empresa',
               'TipoEquipo', 'Marca', 'Modelo', 'Serie', 'Activo', 'EsMovil (Si/No)',
               'CompaniaMovil', 'NumeroMovil', 'SIM', 'EsTemporal (Si/No)', 'FechaDevolucion']

    sample_data = []
    if terminals:
        sample_data.append([
            terminals[0].Name, '12345678-9', 'Juan', 'Perez', 'Operaciones',
            companies[0].Name if companies else '',
            device_types[0].TypeName if device_types else '',
            device_types[0].Brand if device_types else '',
            device_types[0].Model if device_types else '',
            'SAMPLE001', 'Fijo', 'No', '', '', '', 'No', ''
        ])

    df = pd.DataFrame(sample_data if sample_data else [['' for _ in columns]], columns=columns)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Plantilla HGT')

        ws = writer.sheets['Plantilla HGT']

        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='212A37', end_color='212A37', fill_type='solid')
        border = Border(
            left=Side(style='thin', color='212A37'),
            right=Side(style='thin', color='212A37'),
            top=Side(style='thin', color='212A37'),
            bottom=Side(style='thin', color='212A37')
        )

        for col_num, col in enumerate(ws[1], 1):
            col.font = header_font
            col.fill = header_fill
            col.border = border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border

        for col in ws.columns:
            max_length = 10
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 35)

    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Plantilla_CargaMasiva_HGT.xlsx'
    )


@inventory_bp.route('/inventory/bulk_upload', methods=['POST'])
@csrf_required
@limiter.limit("10 per minute")
def bulk_upload():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    if not session.get('permissions', {}).get('bulk_inventory'):
        flash("No tiene permisos para cargar inventario masivo.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    if 'file' not in request.files:
        flash("No se seleccionó ningún archivo.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    file = request.files['file']
    if file.filename == '':
        flash("No se seleccionó ningún archivo.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    valid, err_msg = validate_file_upload(file, ['xlsx', 'xls'])
    if not valid:
        flash(err_msg, "danger")
        return redirect(url_for('inventory.list_inventory'))

    file.stream.seek(0, os.SEEK_END)
    size = file.stream.tell()
    file.stream.seek(0)
    if size > MAX_UPLOAD_SIZE:
        flash(f"El archivo excede el tamaño máximo de {MAX_UPLOAD_SIZE // (1024*1024)} MB.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    try:
        file_bytes = file.read()
        if not file_bytes:
            flash("El archivo está vacío.", "danger")
            return redirect(url_for('inventory.list_inventory'))
        df = pd.read_excel(io.BytesIO(file_bytes))
    except Exception as e:
        flash("Error al leer el archivo. Verifique que sea un Excel válido.", "danger")
        return redirect(url_for('inventory.list_inventory'))

    if df.empty:
        flash("El archivo no contiene filas de datos.", "warning")
        return redirect(url_for('inventory.list_inventory'))

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT TerminalID, Name FROM Terminals")
    terminal_map = {t.Name.strip().lower(): t.TerminalID for t in cursor.fetchall()}

    cursor.execute("SELECT CompanyID, Name FROM Companies")
    company_map = {c.Name.strip().lower(): c.CompanyID for c in cursor.fetchall()}

    cursor.execute("SELECT DeviceTypeID, TypeName, Brand, Model FROM DeviceTypes")
    device_type_map = {}
    for dt in cursor.fetchall():
        key = f"{dt.TypeName}-{dt.Brand}-{dt.Model}".lower()
        device_type_map[key] = dt.DeviceTypeID

    cursor.execute("SELECT StaffID, RUT FROM Staff")
    staff_map = {}
    for s in cursor.fetchall():
        rut_clean = s.RUT.strip().replace('-', '').replace('.', '')
        staff_map[rut_clean] = s.StaffID

    created = 0
    errors = []

    for idx, row in df.iterrows():
        try:
            rut_raw = str(row.get('RUT', '')).strip()
            if not rut_raw or rut_raw == 'nan':
                errors.append(f"Fila {idx+2}: RUT vacío")
                continue

            rut = rut_raw.replace('-', '').replace('.', '')

            terminal_name = str(row.get('Terminal', '')).strip()
            terminal_id = None
            if terminal_name and terminal_name != 'nan':
                terminal_id = terminal_map.get(terminal_name.strip().lower())

            company_name = str(row.get('Empresa', '')).strip()
            company_id = None
            if company_name and company_name != 'nan':
                company_id = company_map.get(company_name.strip().lower())

            device_type_name = str(row.get('TipoEquipo', '')).strip()
            brand = str(row.get('Marca', '')).strip()
            model = str(row.get('Modelo', '')).strip()
            device_key = f"{device_type_name}-{brand}-{model}".lower()
            device_type_id = device_type_map.get(device_key)

            serial = str(row.get('Serie', '')).strip()
            if not serial or serial == 'nan':
                errors.append(f"Fila {idx+2}: Serie vacía")
                continue

            asset_type = str(row.get('Activo', 'Fijo')).strip()
            if asset_type == 'nan':
                asset_type = 'Fijo'

            is_mobile_str = str(row.get('EsMovil (Si/No)', 'No')).strip().lower()
            is_mobile = 1 if is_mobile_str in ['si', 'sí', 'yes', '1'] else 0

            mobile_company = None
            mobile_number = None
            mobile_sim = None
            if is_mobile:
                mobile_company = str(row.get('CompaniaMovil', '')).strip()
                mobile_number = str(row.get('NumeroMovil', '')).strip()
                mobile_sim = str(row.get('SIM', '')).strip()
                if mobile_company == 'nan': mobile_company = None
                if mobile_number == 'nan': mobile_number = None
                if mobile_sim == 'nan': mobile_sim = None

            is_temp_str = str(row.get('EsTemporal (Si/No)', 'No')).strip().lower()
            is_temporary = 1 if is_temp_str in ['si', 'sí', 'yes', '1'] else 0

            expected_return = None
            ret_date_val = row.get('FechaDevolucion')
            if is_temporary and ret_date_val and str(ret_date_val) != 'nan':
                try:
                    expected_return = pd.to_datetime(ret_date_val).to_pydatetime()
                except:
                    pass

            first_name = str(row.get('Nombre', '')).strip()
            last_name = str(row.get('Apellido', '')).strip()
            work_area = str(row.get('AreaTrabajo', '')).strip()
            if first_name == 'nan': first_name = ''
            if last_name == 'nan': last_name = ''
            if work_area == 'nan': work_area = None

            staff_id = staff_map.get(rut)
            if not staff_id and first_name and last_name:
                try:
                    cursor.execute("EXEC sp_CreateStaff @FirstName=?, @LastName=?, @RUT=?, @WorkArea=?, @TerminalID=?",
                                   (first_name, last_name, rut, work_area, terminal_id))
                    conn.commit()

                    cursor.execute("SELECT StaffID FROM Staff WHERE RUT=?", (rut,))
                    result = cursor.fetchone()
                    if result:
                        staff_id = result.StaffID
                        staff_map[rut] = staff_id
                    else:
                        errors.append(f"Fila {idx+2}: No se pudo crear el personal con RUT {rut}")
                        continue
                except Exception as e:
                    errors.append(f"Fila {idx+2}: Error creando personal con RUT {rut}.")
                    continue

            if not device_type_id:
                errors.append(f"Fila {idx+2}: Tipo de equipo no encontrado: {device_type_name} {brand} {model}")
                continue

            try:
                cursor.execute("EXEC sp_CreateAssignment @StaffID=?, @WarehouseID=?, @DeviceTypeID=?, @Serial=?, @AssetType=?, @Date=?, @CompanyID=?, @IsTemporary=?, @ExpectedReturnDate=?, @UserID=?, @IsMobile=?, @MobileCompany=?, @MobileNumber=?, @MobileSIM=?, @TerminalID=?",
                               (staff_id, None, device_type_id, serial, asset_type, datetime.now(), company_id, is_temporary, expected_return, session.get('user_id'), is_mobile, mobile_company, mobile_number, mobile_sim, terminal_id))
                conn.commit()
                created += 1
            except Exception as e:
                err_msg = str(e)
                if isinstance(e, pyodbc.Error) and ("2627" in err_msg or "2601" in err_msg or "UIDX" in err_msg):
                    errors.append(f"Fila {idx+2}: El serial '{serial}' ya existe")
                else:
                    errors.append(f"Fila {idx+2}: {safe_error(err_msg)}")

        except Exception as e:
            errors.append(f"Fila {idx+2}: {safe_error(str(e))}")

    conn.close()

    if created > 0:
        flash(f"Se cargaron {created} equipos correctamente.", "success")

    if errors:
        for err in errors[:5]:
            flash(err, "warning")
        if len(errors) > 5:
            flash(f"...y {len(errors)-5} errores más.", "warning")

    return redirect(url_for('inventory.list_inventory'))
