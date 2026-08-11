from flask import Blueprint, render_template, request, session, redirect, url_for, flash, send_file, current_app
from ..database import get_db_connection
import io
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side

from ..utils.csrf import csrf_required
from ..utils.security import limiter
import pyodbc
from ..utils.security import validate_file_upload, MAX_UPLOAD_SIZE, safe_error

staff_bp = Blueprint('staff', __name__)

@staff_bp.route('/staff', methods=['GET', 'POST'])
@csrf_required
@limiter.limit("20 per minute")
def manage_staff():
    if 'user_id' not in session: return redirect(url_for('auth.login'))

    conn = get_db_connection()
    cursor = conn.cursor()

    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'create':
            if not session.get('permissions', {}).get('create_staff'):
                flash("No tiene permisos para crear personal.", "danger")
                return redirect(url_for('staff.manage_staff'))
            first_name = request.form['first_name']
            last_name = request.form['last_name']
            rut = request.form['rut']
            work_area = request.form['work_area']
            terminal_id = request.form['terminal_id']
            
            try:
                cursor.execute("EXEC sp_CreateStaff @FirstName=?, @LastName=?, @RUT=?, @WorkArea=?, @TerminalID=?", 
                               (first_name, last_name, rut, work_area, terminal_id))
                conn.commit()
                flash("Personal creado exitosamente.", "success")
            except Exception as e:
                err_msg = str(e)
                if "2627" in err_msg or "2601" in err_msg:
                    from flask import current_app
                    current_app.logger.warning(f"Duplicate staff RUT: {rut} from user {session.get('user_id')}")
                    flash("Error: Ese RUT ya existe en el sistema.", "danger")
                else:
                    flash(f"Error al crear el personal.", "danger")
            
        elif action == 'update':
            if not session.get('permissions', {}).get('edit_staff'):
                flash("No tiene permisos para editar personal.", "danger")
                return redirect(url_for('staff.manage_staff'))
            staff_id = request.form['staff_id']
            first_name = request.form['first_name']
            last_name = request.form['last_name']
            rut = request.form['rut']
            work_area = request.form['work_area']
            terminal_id = request.form['terminal_id']
            
            try:
                cursor.execute("EXEC sp_UpdateStaff @StaffID=?, @FirstName=?, @LastName=?, @RUT=?, @WorkArea=?, @TerminalID=?", 
                               (staff_id, first_name, last_name, rut, work_area, terminal_id))
                conn.commit()
                flash("Personal actualizado.", "success")
            except Exception as e:
                err_msg = str(e)
                if "2627" in err_msg or "2601" in err_msg:
                    from flask import current_app
                    current_app.logger.warning(f"Duplicate staff RUT on update: {rut} from user {session.get('user_id')}")
                    flash("Error: No se pudo actualizar. Ese RUT ya está en uso.", "danger")
                else:
                    flash(f"Error al actualizar el personal.", "danger")
            
        elif action == 'delete':
            if not session.get('permissions', {}).get('delete_staff'):
                flash("No tiene permisos para eliminar personal.", "danger")
                return redirect(url_for('staff.manage_staff'))
            staff_id = request.form['staff_id']
            cursor.execute("SELECT COUNT(*) FROM Assignments WHERE StaffID=?", (staff_id,))
            if cursor.fetchone()[0] > 0:
                flash("No se puede eliminar: El funcionario aún tiene equipos asignados a su nombre.", "danger")
            else:
                cursor.execute("EXEC sp_DeleteStaff @StaffID=?", (staff_id,))
                conn.commit()
                flash("Personal eliminado.", "success")
            
        return redirect(url_for('staff.manage_staff'))

    cursor.execute("EXEC sp_GetStaff")
    staff_list = cursor.fetchall()
    
    cursor.execute("EXEC sp_GetTerminals")
    terminals = cursor.fetchall()
    
    conn.close()
    return render_template('staff.html', staff_list=staff_list, terminals=terminals)


@staff_bp.route('/staff/template')
def download_staff_template():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT TerminalID, Name FROM Terminals ORDER BY Name")
    terminals = cursor.fetchall()
    conn.close()

    columns = ['RUT', 'Nombre', 'Apellido', 'AreaTrabajo', 'Terminal']
    
    sample_terminal = terminals[0].Name if terminals else 'Terminal Norte'
    sample_data = [[f'12345678-9', 'Juan', 'Perez', 'Operaciones', sample_terminal]]
    
    df = pd.DataFrame(sample_data, columns=columns)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Plantilla Personal')
        ws = writer.sheets['Plantilla Personal']

        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='212A37', end_color='212A37', fill_type='solid')
        border = Border(
            left=Side(style='thin', color='212A37'),
            right=Side(style='thin', color='212A37'),
            top=Side(style='thin', color='212A37'),
            bottom=Side(style='thin', color='212A37')
        )
        
        for col in ws[1]:
            col.font = header_font
            col.fill = header_fill
            col.border = border
            
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                
        for col in ws.columns:
            max_length = 10
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 35)

    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Plantilla_Personal_HGT.xlsx'
    )


@staff_bp.route('/staff/bulk_upload', methods=['POST'])
@csrf_required
@limiter.limit("10 per minute")
def bulk_upload_staff():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    if not session.get('permissions', {}).get('bulk_staff'):
        flash("No tiene permisos para carga masiva de personal.", "danger")
        return redirect(url_for('staff.manage_staff'))

    if 'file' not in request.files:
        flash("No se seleccionó ningún archivo.", "danger")
        return redirect(url_for('staff.manage_staff'))

    file = request.files['file']
    if file.filename == '':
        flash("No se seleccionó ningún archivo.", "danger")
        return redirect(url_for('staff.manage_staff'))

    valid, err_msg = validate_file_upload(file, ['xlsx', 'xls'])
    if not valid:
        flash(err_msg, "danger")
        return redirect(url_for('staff.manage_staff'))

    file.stream.seek(0, os.SEEK_END)
    size = file.stream.tell()
    file.stream.seek(0)
    if size > MAX_UPLOAD_SIZE:
        flash(f"El archivo excede el tamaño máximo de {MAX_UPLOAD_SIZE // (1024*1024)} MB.", "danger")
        return redirect(url_for('staff.manage_staff'))

    try:
        file_bytes = file.read()
        if not file_bytes:
            flash("El archivo está vacío.", "danger")
            return redirect(url_for('staff.manage_staff'))
        df = pd.read_excel(io.BytesIO(file_bytes))
    except Exception as e:
        flash(f"Error al leer el archivo. Verifique que sea un Excel válido.", "danger")
        return redirect(url_for('staff.manage_staff'))

    if df.empty:
        flash("El archivo no contiene filas de datos.", "warning")
        return redirect(url_for('staff.manage_staff'))

    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT TerminalID, Name FROM Terminals")
    terminal_map = {t.Name.strip().lower(): t.TerminalID for t in cursor.fetchall()}

    cursor.execute("SELECT StaffID, RUT FROM Staff")
    staff_map = {}
    for s in cursor.fetchall():
        rut_clean = s.RUT.strip().replace('-', '').replace('.', '')
        staff_map[rut_clean] = s.StaffID

    created = 0
    errors = []

    for idx, row in df.iterrows():
        try:
            rut_raw = str(row.get('RUT', '')).strip()
            if not rut_raw or rut_raw == 'nan':
                errors.append(f"Fila {idx+2}: RUT vacío")
                continue

            rut = rut_raw.replace('-', '').replace('.', '')

            first_name = str(row.get('Nombre', '')).strip()
            last_name = str(row.get('Apellido', '')).strip()
            work_area = str(row.get('AreaTrabajo', '')).strip()
            
            if first_name == 'nan': first_name = ''
            if last_name == 'nan': last_name = ''
            if work_area == 'nan': work_area = None

            if not first_name or not last_name:
                errors.append(f"Fila {idx+2}: Nombre o Apellido vacío")
                continue

            terminal_name = str(row.get('Terminal', '')).strip()
            terminal_id = None
            if terminal_name and terminal_name != 'nan':
                terminal_id = terminal_map.get(terminal_name.strip().lower())

            if staff_map.get(rut):
                errors.append(f"Fila {idx+2}: El RUT {rut_raw} ya existe")
                continue

            try:
                cursor.execute("EXEC sp_CreateStaff @FirstName=?, @LastName=?, @RUT=?, @WorkArea=?, @TerminalID=?",
                               (first_name, last_name, rut, work_area, terminal_id))
                conn.commit()
                
                cursor.execute("SELECT StaffID FROM Staff WHERE RUT=?", (rut,))
                result = cursor.fetchone()
                if result:
                    staff_map[rut] = result.StaffID
                created += 1
            except Exception as e:
                err_msg = str(e)
                if isinstance(e, pyodbc.Error) and ("2627" in err_msg or "2601" in err_msg):
                    errors.append(f"Fila {idx+2}: El RUT {rut_raw} ya existe en la base de datos")
                else:
                    errors.append(f"Fila {idx+2}: {safe_error(err_msg)}")

        except Exception as e:
            errors.append(f"Fila {idx+2}: {safe_error(str(e))}")

    conn.close()

    if created > 0:
        flash(f"Se cargaron {created} personas correctamente.", "success")

    if errors:
        for err in errors[:5]:
            flash(err, "warning")
        if len(errors) > 5:
            flash(f"...y {len(errors)-5} errores más.", "warning")

    return redirect(url_for('staff.manage_staff'))


@staff_bp.route('/staff/export')
def export_staff():
    if 'user_id' not in session:
        return redirect(url_for('auth.login'))

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("EXEC sp_GetStaff")
    staff_list = cursor.fetchall()
    conn.close()

    df = pd.DataFrame([{
        'ID': s.StaffID,
        'Nombre': s.FirstName,
        'Apellido': s.LastName,
        'RUT': s.RUT,
        'AreaTrabajo': s.WorkArea,
        'Terminal': s.TerminalName
    } for s in staff_list])

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Personal')
        ws = writer.sheets['Personal']

        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='212A37', end_color='212A37', fill_type='solid')
        for col in ws[1]:
            col.font = header_font
            col.fill = header_fill
        for col in ws.columns:
            max_length = 10
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 35)

    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='Personal_HGT.xlsx'
    )
